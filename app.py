import streamlit as st
import pandas as pd
from datetime import datetime
import os
import json
import io
import uuid
from PIL import Image as PILImage
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ==========================================
# 1. Page Configuration
# ==========================================
st.set_page_config(
    page_title="차량 및 장구류 점검 | Data Intel PRO",
    page_icon="🚔",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ==========================================
# 2. Settings & Constants
# ==========================================
DATA_DIR = os.path.expanduser("~/.dataintelligence_pro/inspections")
DB_DIR = os.path.expanduser("~/.dataintelligence_pro/db")
IMAGE_DIR = os.path.expanduser("~/.dataintelligence_pro/images")
os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(DB_DIR, exist_ok=True)
os.makedirs(IMAGE_DIR, exist_ok=True)

BRANCHES = ["중앙", "강북", "서대문", "고양", "의정부", "남양주", "강릉", "원주"]

CHECK_ITEMS = {
    "1. 차량 외부 상태": [
        {"id": "ext_1", "title": "차량 외부 도장 상태"},
        {"id": "ext_2", "title": "차량 외부 파손 상태"},
        {"id": "ext_3", "title": "차량 CI 및 BI 랩핑 상태"},
        {"id": "ext_4", "title": "블랙박스 녹화 및 이상 유무"}
    ],
    "2. 엔진 및 구동 시스템": [
        {"id": "eng_1", "title": "엔진오일 누유 및 잔량"},
        {"id": "eng_2", "title": "냉각수 및 라디에이터"},
        {"id": "eng_3", "title": "팬벨트 및 풀리 상태"},
        {"id": "eng_4", "title": "에어크리너 오염도"},
        {"id": "eng_5", "title": "배터리 전압 및 단자"}
    ],
    "3. 전·후축 및 제동 장치": [
        {"id": "axle_1", "title": "타이어 마모 및 공기압"},
        {"id": "axle_2", "title": "브레이크 라이닝/패드"},
        {"id": "axle_3", "title": "차축 베어링 및 구리스"},
        {"id": "axle_4", "title": "서스펜션/쇼바 상태"},
        {"id": "axle_5", "title": "조향 링크 및 유격"}
    ],
    "4. 유압 및 실린더 시스템": [
        {"id": "hydro_1", "title": "유압유 누유 및 잔량"},
        {"id": "hydro_2", "title": "컨트롤 밸브 작동"},
        {"id": "hydro_3", "title": "유압 호스 손상 여부"},
        {"id": "hydro_4", "title": "실린더 로드 부식/누유"},
        {"id": "hydro_5", "title": "작업 장치 핀 유격"}
    ],
    "5. 전기장치 및 안전": [
        {"id": "elec_1", "title": "계기판 경고등 및 게이지"},
        {"id": "elec_2", "title": "전·후방 등화 장치"},
        {"id": "elec_3", "title": "후진 알람 및 경음기"},
        {"id": "elec_4", "title": "와이퍼 및 워셔액"},
        {"id": "elec_5", "title": "소화기 및 안전장구"}
    ]
}

# ==========================================
# 3. Custom CSS Injection
# ==========================================
def load_css():
    css_file = os.path.join(os.path.dirname(__file__), "ui", "style.css")
    if os.path.exists(css_file):
        with open(css_file, "r", encoding="utf-8") as f:
            st.markdown(f'<style>{f.read()}</style>', unsafe_allow_html=True)
load_css()

# ==========================================
# 4. Helper Functions
# ==========================================
@st.dialog("🔍 상세점검내역 가이드", width="large")
def show_inspection_details(item_title, item_detail):
    st.markdown(f"### {item_title}")
    st.info("아래의 점검 가이드라인에 따라 꼼꼼하게 확인해 주십시오.")
    for line in item_detail.split('\n'):
        st.markdown(f"- **{line}**")
    st.markdown("<br>", unsafe_allow_html=True)
    if st.button("확인 완료 (닫기)", type="primary"):
        st.rerun()

def get_vehicle_db_df():
    db_path = os.path.join(DB_DIR, "vehicles.csv")
    
    # 클라우드 배포 환경(Streamlit Cloud 등)에서 DB 파일이 없는 경우,
    # 프로젝트 내부에 번들링된 초기 DB 파일(vehicles_init.csv)로 자동 초기화합니다.
    if not os.path.exists(db_path):
        init_db_path = os.path.join(os.path.dirname(__file__), "vehicles_init.csv")
        if os.path.exists(init_db_path):
            import shutil
            shutil.copy(init_db_path, db_path)

    if os.path.exists(db_path):
        df = pd.read_csv(db_path)
        # 하위 호환성을 위해 컬럼 확인 (BOM 및 공백 제거)
        df.rename(columns=lambda x: x.replace('\ufeff', '').strip(), inplace=True)
        for col in ["지사", "차량번호", "구역번호"]:
            if col not in df.columns:
                df[col] = ""
        df["지사"] = df["지사"].astype(str).str.strip()
        return df
    return pd.DataFrame(columns=["지사", "차량번호", "구역번호"])

def load_all_inspections():
    data = []
    for f in os.listdir(DATA_DIR):
        if f.endswith(".json"):
            with open(os.path.join(DATA_DIR, f), "r", encoding="utf-8") as file:
                try:
                    record = json.load(file)
                    record["_filename"] = f
                    data.append(record)
                except:
                    pass
    return data

def save_image_bytes(img_bytes, filename_prefix="img", ext="jpg"):
    filename = f"{filename_prefix}_{uuid.uuid4().hex[:8]}.{ext}"
    filepath = os.path.join(IMAGE_DIR, filename)
    with open(filepath, "wb") as f:
        f.write(img_bytes)
    return filepath

# ==========================================
# 5. Sidebar Navigation
# ==========================================
with st.sidebar:
    st.markdown("## 🧭 네비게이션")
    page = st.radio("이동할 페이지 선택", ["📋 현장 점검 입력", "📊 관리자 대시보드"])
    st.divider()
    st.markdown("### 🚀 배포 및 엑셀 기능")
    st.info("관리자 탭에서 점검 데이터를 엑셀로 다운로드하고 차량 DB를 업로드할 수 있습니다.")

# ==========================================
# PAGE 1: USER INSPECTION FORM
# ==========================================
if page == "📋 현장 점검 입력":
    # 5.1 Initialize States
    if "temp_photos" not in st.session_state:
        st.session_state["temp_photos"] = {}

    for category, items in CHECK_ITEMS.items():
        for item in items:
            if item["id"] not in st.session_state:
                st.session_state[item["id"]] = None
            if f"memo_{item['id']}" not in st.session_state:
                st.session_state[f"memo_{item['id']}"] = ""

    today_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    st.markdown(f"<div class='nav-bar-custom'><h1>장비 종합 정밀 점검 시스템</h1><span>{today_str}</span></div>", unsafe_allow_html=True)

    # 전역 대시보드 지표 (전체 누적 데이터 기반)
    inspections = load_all_inspections()
    total_vehicles = len(inspections)
    ng_vehicles = sum(1 for d in inspections if d.get("ng_count", 0) > 0 or d.get("warn_count", 0) > 0)
    ok_vehicles = total_vehicles - ng_vehicles

    # 진행률 표시
    total_items = sum(len(items) for items in CHECK_ITEMS.values())
    checked_items = sum(1 for cat, items in CHECK_ITEMS.items() for item in items if st.session_state.get(item["id"]) is not None)
    progress_percent = int((checked_items / total_items) * 100) if total_items > 0 else 0

    st.markdown(f"<div style='font-size: 0.95rem; margin-top: 10px; margin-bottom: 5px; font-weight: 700; color: #64748b;'>점검 진행률: <span style='color: #2563eb;'>{progress_percent}%</span> ({checked_items}/{total_items})</div>", unsafe_allow_html=True)
    st.progress(progress_percent / 100)
    st.markdown("<br>", unsafe_allow_html=True)

    m1, m2, m3 = st.columns(3)
    m1.metric("총 점검 완료 차량", f"{total_vehicles}대")
    m2.metric("주의 및 정비 필요", f"{ng_vehicles}대", f"-조치 요망" if ng_vehicles > 0 else None, delta_color="inverse")
    m3.metric("상태 양호", f"{ok_vehicles}대", "전원 양호" if ok_vehicles == total_vehicles and total_vehicles > 0 else None)

    with st.container(border=True):
        st.markdown("### 📋 현장 기본 정보")
        c1, c2, c3 = st.columns(3)
        df_db = get_vehicle_db_df()
        
        with c1:
            branch = st.selectbox("방문 지사", options=BRANCHES)
        with c2:
            if not df_db.empty:
                filtered_df = df_db[df_db["지사"] == branch]
                if filtered_df.empty:
                    car_num = st.selectbox("차량 번호 선택", options=["차량 없음"])
                    zone_num = ""
                else:
                    car_num = st.selectbox("차량 번호 선택", options=filtered_df["차량번호"].tolist())
                    zone_num = filtered_df[filtered_df["차량번호"] == car_num]["구역번호"].values[0]
            else:
                st.warning("⚠️ 차량 DB가 비어있습니다.")
                car_num = None
                zone_num = ""
        with c3:
            st.text_input("구역 번호", value=zone_num, disabled=True)
            
        mileage = st.number_input("누적 주행거리 (km)", value=0, step=1)

    # 5.2 Responsive Table Checklist
    for idx_cat, (category, items) in enumerate(CHECK_ITEMS.items()):
        # 첫 번째 카테고리만 기본으로 펼쳐두기
        is_expanded = (idx_cat == 0)
        with st.expander(f"📋 {category}", expanded=is_expanded):
            # Header
            h_col1, h_col2, h_col3 = st.columns([2, 1.5, 2])
            h_col1.markdown("<div style='color: #64748b; font-size: 0.9rem; font-weight: 700; padding-bottom: 5px;'>점검 항목</div>", unsafe_allow_html=True)
            h_col2.markdown("<div style='color: #64748b; font-size: 0.9rem; font-weight: 700; padding-bottom: 5px;'>상태</div>", unsafe_allow_html=True)
            h_col3.markdown("<div style='color: #64748b; font-size: 0.9rem; font-weight: 700; padding-bottom: 5px;'>비고</div>", unsafe_allow_html=True)
            st.divider()

            for idx, item in enumerate(items):
                i_col1, i_col2, i_col3 = st.columns([2, 1.5, 2])
                with i_col1:
                    st.markdown(f"<div style='padding-top: 15px; font-weight: 600; color: #1e293b;'>{item['title']}</div>", unsafe_allow_html=True)
                with i_col2:
                    st.radio("상태", ["양호", "주의", "불량"], key=item["id"], horizontal=True, label_visibility="collapsed", index=None)
                with i_col3:
                    st.text_input("비고", key=f"memo_{item['id']}", placeholder="이상 내용 입력", label_visibility="collapsed")
                
                if idx < len(items) - 1:
                    st.markdown("<hr style='margin: 0px; opacity: 0.3;'>", unsafe_allow_html=True)

    # 5.3 Photo Capture & Temp Storage
    with st.container(border=True):
        st.markdown("### 📸 현장 증빙 사진 (최대 10장)")
        st.info("카메라 촬영이나 파일 업로드 시 임시 갤러리에 계속 누적됩니다. 잘못 찍은 사진은 [X 삭제] 버튼으로 지울 수 있습니다.")
        
        cam_col, up_col = st.columns(2)
        with cam_col:
            cam_photo = st.camera_input("실시간 촬영", label_visibility="collapsed")
            if cam_photo:
                # 고유 해시 대신 name/size 조합으로 id 생성 (Streamlit 호환성)
                fid = f"{cam_photo.name}_{cam_photo.size}"
                if fid not in st.session_state["temp_photos"]:
                    if len(st.session_state["temp_photos"]) >= 10:
                        st.error("⚠️ 사진은 최대 10장까지만 업로드 가능합니다.")
                    else:
                        st.session_state["temp_photos"][fid] = {"name": cam_photo.name, "bytes": cam_photo.getvalue()}
                        st.rerun()

        with up_col:
            uploaded_files = st.file_uploader("파일 업로드", accept_multiple_files=True, type=['png', 'jpg', 'jpeg'], key="uploader")
            if uploaded_files:
                added = False
                for f in uploaded_files:
                    fid = f"{f.name}_{f.size}"
                    if fid not in st.session_state["temp_photos"]:
                        if len(st.session_state["temp_photos"]) >= 10:
                            st.error("⚠️ 더 이상 추가할 수 없습니다. (최대 10장)")
                            break
                        else:
                            st.session_state["temp_photos"][fid] = {"name": f.name, "bytes": f.getvalue()}
                            added = True
                if added:
                    st.rerun()

        if st.session_state["temp_photos"]:
            st.markdown("#### 📥 임시 갤러리 (제출 시 최종 저장)")
            st.progress(len(st.session_state["temp_photos"]) / 10, text=f"{len(st.session_state['temp_photos'])} / 10 장 임시저장됨")
            
            p_cols = st.columns(5)
            keys_to_delete = []
            
            for idx, (fid, pdata) in enumerate(st.session_state["temp_photos"].items()):
                with p_cols[idx % 5]:
                    st.image(pdata["bytes"], use_container_width=True)
                    if st.button("❌ 삭제", key=f"del_{fid}", use_container_width=True):
                        keys_to_delete.append(fid)
            
            if keys_to_delete:
                for k in keys_to_delete:
                    del st.session_state["temp_photos"][k]
                st.rerun()

    with st.container(border=True):
        st.markdown("### 📝 종합 의견")
        memo = st.text_area("특이사항 기록", placeholder="차량 외관 스크래치 등 특이사항을 입력해주세요.", height=100)

    # 5.4 Submit
    st.markdown("<br>", unsafe_allow_html=True)
    if st.button("🚀 점검 완료 및 데이터 전송", use_container_width=True, type="primary"):
        if checked_items < total_items:
            st.error("⚠️ 미점검 항목이 있습니다. 모든 항목을 점검해주세요.")
        elif not car_num or not car_num.strip() or car_num == "차량 없음":
            st.error("⚠️ 차량 번호를 정확히 선택해주세요.")
        elif len(st.session_state["temp_photos"]) > 10:
            st.error("⚠️ 사진 개수를 10장 이하로 줄여주세요.")
        else:
            current_ok, current_warn, current_ng = 0, 0, 0
            for cat, items in CHECK_ITEMS.items():
                for it in items:
                    val = st.session_state.get(it["id"])
                    if val == "양호": current_ok += 1
                    elif val == "주의": current_warn += 1
                    elif val == "불량": current_ng += 1
            current_total = current_ok + current_warn + current_ng

            saved_image_paths = []
            for fid, pdata in st.session_state["temp_photos"].items():
                ext = pdata["name"].split(".")[-1] if "." in pdata["name"] else "jpg"
                saved_path = save_image_bytes(pdata["bytes"], filename_prefix=f"{car_num}_img", ext=ext)
                saved_image_paths.append(saved_path)

            report_data = {
                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "car_num": car_num,
                "branch": branch,
                "zone_num": zone_num,
                "mileage": mileage,
                "total_count": current_total,
                "ok_count": current_ok,
                "warn_count": current_warn,
                "ng_count": current_ng,
                "memo": memo,
                "images": saved_image_paths
            }
            
            for cat, items in CHECK_ITEMS.items():
                for item in items:
                    report_data[item["title"]] = st.session_state[item["id"]]
                    report_data[f"memo_{item['title']}"] = st.session_state[f"memo_{item['id']}"]
            
            save_path = os.path.join(DATA_DIR, f"report_{car_num}_{datetime.now().strftime('%Y%m%d%H%M%S')}.json")
            with open(save_path, "w", encoding="utf-8") as f:
                json.dump(report_data, f, ensure_ascii=False, indent=4)
                
            # 제출 후 임시 상태 초기화
            st.session_state["temp_photos"].clear()
            for cat, items in CHECK_ITEMS.items():
                for item in items:
                    st.session_state[item["id"]] = None
                    st.session_state[f"memo_{item['id']}"] = ""
            

            st.success(f"✅ {car_num} 점검 데이터 및 사진({len(saved_image_paths)}장)이 성공적으로 최종 저장되었습니다!")
            st.balloons()


# ==========================================
# PAGE 2: ADMIN DASHBOARD
# ==========================================
elif page == "📊 관리자 대시보드":
    st.markdown("## 🔐 관리자 인증")
    if "admin_auth" not in st.session_state:
        st.session_state["admin_auth"] = False
        
    if not st.session_state["admin_auth"]:
        pwd = st.text_input("관리자 비밀번호를 입력하세요", type="password")
        if st.button("로그인"):
            if pwd == "admin1234!!":
                st.session_state["admin_auth"] = True
                st.rerun()
            else:
                st.error("비밀번호가 일치하지 않습니다.")
    else:
        st.markdown("<div style='text-align: right;'><small>✅ 관리자 인증 완료</small></div>", unsafe_allow_html=True)
        if st.button("로그아웃"):
            st.session_state["admin_auth"] = False
            st.rerun()
            
        st.divider()
        
        # 차량 DB 관리 패널
        with st.expander("🛠️ 차량 번호 관리 (단일 등록 및 일괄 업로드)"):
            c1, c2 = st.columns(2)
            with c1:
                st.markdown("#### 단일 차량 등록")
                new_branch = st.selectbox("지사 선택", options=BRANCHES, key="new_branch")
                new_car = st.text_input("새로운 차량 번호 입력 (예: 123가4567)")
                new_zone = st.text_input("구역 번호 입력 (예: G0001234)")
                if st.button("추가하기"):
                    if new_car and new_zone:
                        df_db = get_vehicle_db_df()
                        if new_car not in df_db["차량번호"].values:
                            df_db.loc[len(df_db)] = [new_branch, new_car, new_zone]
                            df_db.to_csv(os.path.join(DB_DIR, "vehicles.csv"), index=False)
                            st.success(f"✅ {new_car} 차량이 추가되었습니다.")
                        else:
                            st.warning("이미 존재하는 차량입니다.")
                    else:
                        st.warning("차량 번호와 구역 번호를 모두 입력해주세요.")
            with c2:
                st.markdown("#### 엑셀/CSV 일괄 업로드")
                db_file = st.file_uploader("'지사', '차량번호', '구역번호' 컬럼을 가진 파일 선택", type=['csv', 'xlsx'])
                if db_file:
                    try:
                        df_db = pd.read_csv(db_file) if db_file.name.endswith('.csv') else pd.read_excel(db_file)
                        req_cols = ["지사", "차량번호", "구역번호"]
                        if all(col in df_db.columns for col in req_cols):
                            df_db[req_cols].to_csv(os.path.join(DB_DIR, "vehicles.csv"), index=False)
                            st.success(f"✅ 차량 DB가 업데이트 되었습니다. (총 {len(df_db)}대)")
                        else:
                            st.error(f"파일에 필수 컬럼({', '.join(req_cols)})이 모두 포함되어야 합니다.")
                    except Exception as e:
                        st.error(f"오류 발생: {e}")
                        
        st.divider()
        
        inspections = load_all_inspections()
        if not inspections:
            st.warning("아직 등록된 점검 데이터가 없습니다.")
        else:
            df = pd.DataFrame(inspections)
            df["삭제 선택"] = False
            cols_to_show = ["삭제 선택", "timestamp", "branch", "car_num", "zone_num", "ng_count", "memo", "_filename", "images"]
            
            # 하위 호환성 (과거 V2 데이터에 images 컬럼이 없을 경우)
            for col in cols_to_show:
                if col not in df.columns:
                    df[col] = None
            
            st.markdown("### 🗑️ 점검 내역 관리 및 삭제")
            st.info("삭제할 항목을 체크한 뒤 아래 버튼을 누르면 영구 삭제됩니다. (관련된 사진 파일도 함께 삭제됩니다.)")
            
            edited_df = st.data_editor(
                df[cols_to_show],
                hide_index=True,
                column_config={
                    "삭제 선택": st.column_config.CheckboxColumn("삭제 선택", default=False),
                    "_filename": None, 
                    "images": None 
                },
                disabled=["timestamp", "branch", "car_num", "ng_count", "memo"]
            )
            
            if st.button("🗑️ 선택 항목 영구 삭제", type="primary"):
                to_delete = edited_df[edited_df["삭제 선택"] == True]
                if len(to_delete) > 0:
                    for _, row in to_delete.iterrows():
                        imgs = row.get("images", [])
                        if isinstance(imgs, list):
                            for img_path in imgs:
                                if img_path and os.path.exists(img_path):
                                    os.remove(img_path)
                        json_path = os.path.join(DATA_DIR, row["_filename"])
                        if os.path.exists(json_path):
                            os.remove(json_path)
                    st.success(f"✅ {len(to_delete)}건의 데이터가 영구 삭제되었습니다.")
                    st.rerun()
                else:
                    st.warning("삭제할 항목을 체크해주세요.")

            st.divider()
            
            st.markdown("### 📊 전문가용 사진 통합 엑셀 다운로드")
            st.info("클릭 시 각 차량별 점검 내역과 촬영된 사진(최대 10장)이 가로로 나열되어 포함된 엑셀 보고서가 생성됩니다.")
            
            if st.button("📥 전문 엑셀(사진 포함) 다운로드 생성하기"):
                with st.spinner("엑셀에 사진을 렌더링하는 중입니다... (약간의 시간이 소요될 수 있습니다)"):
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "점검 보고서"
                    
                    headers = ["점검일시", "지사", "차량번호", "구역번호", "경고(주의)건수", "불량건수", "종합특이사항"]
                    for cat, items in CHECK_ITEMS.items():
                        for it in items:
                            headers.append(it["title"] + " (상태)")
                            headers.append(it["title"] + " (비고)")
                    
                    for i in range(1, 11):
                        headers.append(f"사진 {i}")
                    
                    ws.append(headers)
                    
                    header_fill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
                    header_font = Font(color="FFFFFF", bold=True)
                    for col_num, header in enumerate(headers, 1):
                        cell = ws.cell(row=1, column=col_num)
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        
                        if "사진" in header:
                            ws.column_dimensions[get_column_letter(col_num)].width = 28
                        else:
                            ws.column_dimensions[get_column_letter(col_num)].width = 18

                    for r_idx, row_data in enumerate(inspections, start=2):
                        base_data = [
                            row_data.get("timestamp", ""),
                            row_data.get("branch", ""),
                            row_data.get("car_num", ""),
                            row_data.get("zone_num", ""),
                            row_data.get("warn_count", 0),
                            row_data.get("ng_count", 0),
                            row_data.get("memo", "")
                        ]
                        for cat, items in CHECK_ITEMS.items():
                            for it in items:
                                base_data.append(row_data.get(it["title"], "미점검"))
                                base_data.append(row_data.get(f"memo_{it['title']}", ""))
                        
                        for c_idx, val in enumerate(base_data, start=1):
                            cell = ws.cell(row=r_idx, column=c_idx, value=val)
                            cell.alignment = Alignment(vertical="center")
                        
                        ws.row_dimensions[r_idx].height = 130 
                        
                        imgs = row_data.get("images", [])
                        if isinstance(imgs, list):
                            for img_idx, img_path in enumerate(imgs[:10]):
                                if img_path and os.path.exists(img_path):
                                    try:
                                        with PILImage.open(img_path) as img:
                                            # 고해상도 리사이징 (LANCZOS, 5:4 비율)
                                            img_resized = img.resize((180, 144), PILImage.Resampling.LANCZOS)
                                            img_io = io.BytesIO()
                                            img_resized.save(img_io, format="JPEG")
                                            img_io.seek(0)
                                            
                                            xl_img = OpenpyxlImage(img_io)
                                            img_col = len(base_data) + img_idx + 1
                                            col_letter = get_column_letter(img_col)
                                            xl_img.anchor = f"{col_letter}{r_idx}"
                                            xl_img.width = 180
                                            xl_img.height = 144
                                            ws.add_image(xl_img)
                                    except Exception as e:
                                        pass
                    
                    output = io.BytesIO()
                    wb.save(output)
                    excel_data = output.getvalue()
                
                st.download_button(
                    label="💾 보고서 엑셀 저장하기",
                    data=excel_data,
                    file_name=f"사진통합_차량점검보고서_{datetime.now().strftime('%Y%m%d')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary"
                )
